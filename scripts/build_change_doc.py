"""Build a concise Word record from the verified project state."""

from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.8)
for name, size in [
    ("Normal", 11),
    ("Title", 24),
    ("Subtitle", 11),
    ("Heading 1", 16),
    ("Heading 2", 12),
]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = Pt(16 if name == "Normal" else size * 1.25)
    style.paragraph_format.space_before = Pt(8 if name.startswith("Heading") else 0)
    if name == "Heading 2":
        style.paragraph_format.space_after = Pt(3)
doc.core_properties.title = "MYATRI 修改记录与 AstrBot 接入说明"
doc.core_properties.author = "MYATRI"
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run("MYATRI 项目记录  |  ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

title = doc.add_paragraph("MYATRI 修改记录与\nAstrBot 接入说明", "Title")
for style in doc.styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
title_borders = OxmlElement("w:pBdr")
for side in ["top", "bottom", "left", "right", "between"]:
    edge = OxmlElement("w:" + side)
    edge.set(qn("w:val"), "nil")
    title_borders.append(edge)
title._p.get_or_add_pPr().append(title_borders)
doc.add_paragraph("2026年10月21日  主项目 0.1.7  桥接项目 0.1.0", "Subtitle")
doc.add_paragraph(
    "小温暖已从命令行角色发展为可持久保存记忆、按轮次维护的本地 Agent。本记录汇总既有修改及新建 AstrBot 项目的边界，便于后续启动、恢复和继续开发。微信群真实连接仍需配置并实测平台适配器。"
)
doc.add_heading("当前配置", level=1)
rows = [
    ("私聊入口与库", "启动小温暖.cmd；var/xiaowennuan，内部 agent ID 为 atri"),
    ("当前人格", "角色卡1.2；最近人工版本 persona@5，后续版本可继续增长"),
    ("日预算", "50M tokens；按运行库账本统计，不代表账号所有调用总额"),
    ("上下文与压缩", "1M配置上限；650k自动压缩软阈值，压缩前处理依赖任务"),
    ("完整诊断", "每30个成功对话轮次"),
    ("记忆与情绪整理", "每50轮；临时状态与事件抽取合并维护"),
    ("深度整理与人格审阅", "每3批约150轮反思；每6批约300轮人格审阅"),
    ("投影与后台轮询", "约60次提交刷新；检查点及退出刷新；空队列5秒轮询"),
    ("日常回复", "40—80字目标，80字本地上限；短回复不硬凑字"),
    ("前台输出", "最多512 tokens；80字裁剪只作用于最终聊天正文"),
    ("时间与原话", "Asia/Shanghai；原话即时落盘，时间与正文分开显示"),
]
table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width = Inches(1.65)
table.columns[1].width = Inches(5.15)
for cell, text in zip(table.rows[0].cells, ["项目", "当前值"]):
    cell.text = text
for row in rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, row):
        cell.text = text
for i, row in enumerate(table.rows):
    for j, cell in enumerate(row.cells):
        cell.width = Inches(1.65 if j == 0 else 5.15)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        properties = cell._tc.get_or_add_tcPr()
        margins = OxmlElement("w:tcMar")
        for side in ["top", "bottom", "left", "right"]:
            item = OxmlElement("w:" + side)
            item.set(qn("w:w"), "90")
            item.set(qn("w:type"), "dxa")
            margins.append(item)
        properties.append(margins)
        border = OxmlElement("w:tcBorders")
        for side in ["top", "bottom", "left", "right"]:
            item = OxmlElement("w:" + side)
            for attr, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                item.set(qn("w:" + attr), value)
            border.append(item)
        properties.append(border)
        shade = OxmlElement("w:shd")
        shade.set(
            qn("w:fill"), "24445C" if i == 0 else ("F2F5F7" if i % 2 == 0 else "FFFFFF")
        )
        properties.append(shade)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                run.font.bold = i == 0
                if i == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
header = OxmlElement("w:tblHeader")
table.rows[0]._tr.get_or_add_trPr().append(header)
note = doc.add_paragraph(
    "50轮、150轮和300轮是维护节奏，计数达到阈值本身不构成人格变化证据。650k真实长上下文及长期语义表现仍需验证。"
)
note.paragraph_format.space_before = Pt(8)

doc.add_page_break()
doc.add_heading("既有修改记录", level=1)
sections = [
    (
        "一键启动与凭据",
        "新增CMD、PowerShell和Python启动器，自动检查环境与既有库，使用Windows凭据读取模型Key。hell.exe启动错误通过纯ASCII CMD、CRLF换行和显式PowerShell路径处理；中文名称由PowerShell预设生成。",
    ),
    (
        "ATRI记忆继承与身份修复",
        "备份后完整恢复ATRI历史，迁移时保留61条原始消息和413个既有revision，内部atri域不改。后续将继承消息作为带来源的历史数据，首尾加入当前身份锚点，避免旧自述覆盖小温暖。新旧库后续不自动同步。",
    ),
    (
        "角色卡与版本化更新",
        "1.0原文逐字节保留。1.1调整为轻柔、文雅、灵动、含蓄俏皮；1.2减少反问、征求意见和邮件式总结。当前最近人工版本为persona@5。所有修改经过统一writer与版本检查，保留继承边界、历史人格和原话。",
    ),
    (
        "调用频率与本地开销",
        "逐消息认知改为成功轮次批处理，当前采用30/50轮及3/6批倍率；状态激活改为本地选择。压缩前先处理依赖材料。加入档案缓存、深拷贝和增量投影，完整协议按检查点保存；原话、未完成协议及预算仍即时持久化。",
    ),
    (
        "情绪与短聊天",
        "PAD情绪、慢心境、事件评价和临时余波支持撤回，不自动改核心人格。40—80字通过生成指令约束，超长本地按句或分句收住，不增加LLM改写；不足40字不填充套话，工具协议不裁剪。",
    ),
    (
        "消息时间与终端显示",
        "模型上下文包含每条消息的时区时间，事件发生时间与发送时间区分。终端分开显示你、小温暖和系统，去除正文复述的前导时间标签。代码更新后需退出旧窗口再启动；不会自动热加载。",
    ),
]
for heading, text in sections:
    doc.add_heading(heading, level=2)
    doc.add_paragraph(text)
doc.add_heading("验证与恢复", level=2)
doc.add_paragraph(
    "0.1.7完成245项完整离线回归、29项定向检查和独立wheel安装smoke。最近短聊天更新前已有105条原话及490个历史revision，均核对未改；数字为当时快照。更新前备份保存在var/backups。仍开放部分检索、撤回与长期语义问题，测试通过不等于无限期记忆验收。"
)

doc.add_page_break()
doc.add_heading("AstrBot 独立项目", level=1)
doc.add_paragraph(
    "新项目位于 D:/DS_workspace/MYATRI-AstrBot。已克隆AstrBot 4.29.1，commit固定为f2404c191716a64daaa422be7f3c567cace50d39。新项目使用冻结的MYATRI 0.1.7 wheel，依赖和虚拟环境独立；没有复制原私聊运行库或模型Key。"
)
doc.add_heading("本次新增", level=2)
doc.add_paragraph(
    "AstrBot负责群消息事件和平台收发，本机HTTP后端负责白名单、去重、群隔离及调用小温暖。入站信封保存群ID、成员ID、昵称、原平台时间和原文；同群共享上下文，跨群与私聊分开。成员稳定ID必须由渠道可靠提供，昵称不是身份认证。"
)
doc.add_paragraph(
    "投递分为prepared、sending、sent与unknown。插件先原子领取再发送，结果未知时不自动重发；sent仅表示发送接口成功返回，不证明成员已读。所有桥接群共享50M日预算，原命令行私聊另计。现有E01主动社交仍未实现。"
)
doc.add_heading("启动与接入顺序", level=2)
for text in [
    "1 运行Start-Bridge.cmd，默认启动本机offline后端并准备插件。",
    "2 运行Start-AstrBot.cmd，进入管理面板，选择实际支持普通微信群的适配器。",
    "3 核对电脑微信版本，以及稳定群ID、成员ID、消息ID和@组件。",
    "4 在后端与插件中配置相同群白名单，先用测试群验证offline回复。",
    "5 再按需要切换后端live模式。专用AstrBot关闭内置LLM自动回复，避免双重调用。",
]:
    doc.add_paragraph(text)
doc.add_heading("已验证与尚未验证", level=2)
doc.add_paragraph(
    "14项离线检查通过，包含实际AstrBot事件对象、注册接口及插件，经本机HTTP调用真实MYATRI引擎，发送使用本地捕获器。覆盖重复消息、跨群隔离、原文时间、并发领取、unknown不重发和人格更新后重启。当前offline、空白名单、插件关闭；没有微信登录、真实群消息或真实LLM调用。"
)
doc.add_paragraph(
    "普通微信群连接层尚未选定。官方个人微信接口不能直接视作群聊已接通；应先按实际微信版本验证适配器。私聊记忆若要向群共享，需另行选择可公开材料与可见范围。"
)
doc.add_heading("完整记录与来源", level=2)
doc.add_paragraph(
    "完整修改历史：原项目docs/项目修改记录与当前配置_2026-10-21.md；新项目docs/MYATRI既有修改记录.md。启动、接口限制及验收见新项目README.md、docs/接入设计与验收.md和docs/verification.json。"
)
doc.add_paragraph(
    "上游：https://github.com/AstrBotDevs/AstrBot\n插件文档：https://docs.astrbot.app/dev/star/plugin-new.html\n微信群需求：https://github.com/Tencent/openclaw-weixin/issues/327"
)

output = ROOT / "docs/MYATRI修改记录与AstrBot接入说明.docx"
doc.save(output)
shutil.copyfile(output, Path("D:/DS_workspace/MYATRI/docs") / output.name)
print(output)
